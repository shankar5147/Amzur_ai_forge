from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass
from pathlib import Path

import anyio
import openpyxl
from pydantic import BaseModel
from pypdf import PdfReader

from app.models.db_models import Attachment
from app.services.file_storage_service import FileStorageService

try:
    import docx as python_docx
except ImportError:  # pragma: no cover
    python_docx = None

try:
    from pptx import Presentation as PptxPresentation
except ImportError:  # pragma: no cover
    PptxPresentation = None


class AttachmentPreviewError(Exception):
    pass


@dataclass(frozen=True)
class AttachmentPreview:
    preview_type: str
    columns: list[str]
    rows: list[list[str]]
    content: str | None
    truncated: bool


class AttachmentPreviewService:
    _max_rows = 20
    _max_cols = 12
    _max_text_chars = 8000

    def __init__(self, storage: FileStorageService | None = None) -> None:
        self._storage = storage or FileStorageService()

    async def build_preview(self, attachment: Attachment) -> AttachmentPreview:
        path = self._storage.resolve_path(attachment.file_path)
        if not path.exists():
            raise AttachmentPreviewError("Attachment file was not found on disk.")

        mime = attachment.mime_type

        if mime in {"text/csv", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}:
            if mime == "text/csv":
                columns, rows = await anyio.to_thread.run_sync(self._preview_csv, path)
            else:
                columns, rows = await anyio.to_thread.run_sync(self._preview_xlsx, path)
            return AttachmentPreview(
                preview_type="table",
                columns=columns,
                rows=rows,
                content=None,
                truncated=len(rows) >= self._max_rows,
            )

        if mime == "application/pdf":
            content = await anyio.to_thread.run_sync(self._extract_pdf, path)
            content, truncated = self._truncate(content)
            return AttachmentPreview("document", [], [], content, truncated)

        if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            content = await anyio.to_thread.run_sync(self._extract_docx, path)
            content, truncated = self._truncate(content)
            return AttachmentPreview("document", [], [], content, truncated)

        if mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
            content = await anyio.to_thread.run_sync(self._extract_pptx, path)
            content, truncated = self._truncate(content)
            return AttachmentPreview("document", [], [], content, truncated)

        if mime in {
            "text/plain",
            "text/markdown",
            "text/html",
            "text/xml",
            "application/xml",
            "application/json",
            "application/rtf",
        }:
            content = await anyio.to_thread.run_sync(path.read_text, "utf-8", "ignore")
            if mime == "application/json":
                try:
                    content = json.dumps(json.loads(content), indent=2)
                except json.JSONDecodeError:
                    pass
            content, truncated = self._truncate(content)
            return AttachmentPreview("text", [], [], content, truncated)

        # Legacy binary office formats are allowed to upload but not parsed for inline preview.
        if mime in {"application/msword", "application/vnd.ms-powerpoint", "application/vnd.ms-excel"}:
            return AttachmentPreview(
                preview_type="document",
                columns=[],
                rows=[],
                content="Inline preview is not available for this legacy Office format. Please open/download the file.",
                truncated=False,
            )

        return AttachmentPreview(
            preview_type="unsupported",
            columns=[],
            rows=[],
            content="Inline preview is not available for this file type.",
            truncated=False,
        )

    def _preview_csv(self, path: Path) -> tuple[list[str], list[list[str]]]:
        with path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
            reader = csv.reader(handle)
            rows = list(reader)

        if not rows:
            return [], []

        header = [str(cell) for cell in rows[0]][: self._max_cols]
        body = [[str(cell) for cell in row][: self._max_cols] for row in rows[1 : self._max_rows + 1]]
        return header, body

    def _preview_xlsx(self, path: Path) -> tuple[list[str], list[list[str]]]:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            values = list(sheet.iter_rows(min_row=1, max_row=self._max_rows + 1, values_only=True))
        finally:
            workbook.close()

        if not values:
            return [], []

        header = [self._cell_to_string(value) for value in values[0]][: self._max_cols]
        rows = [
            [self._cell_to_string(value) for value in row][: self._max_cols]
            for row in values[1 : self._max_rows + 1]
        ]
        return header, rows

    @staticmethod
    def _extract_pdf(path: Path) -> str:
        reader = PdfReader(str(path))
        chunks: list[str] = []
        for page in reader.pages:
            chunks.append(page.extract_text() or "")
        return "\n".join(chunks).strip() or "No extractable text found in PDF."

    @staticmethod
    def _extract_docx(path: Path) -> str:
        if python_docx is None:
            return "DOCX parser dependency is not installed."
        doc = python_docx.Document(str(path))
        lines = [para.text.strip() for para in doc.paragraphs if para.text and para.text.strip()]
        return "\n".join(lines).strip() or "No extractable text found in DOCX."

    @staticmethod
    def _extract_pptx(path: Path) -> str:
        if PptxPresentation is None:
            return "PPTX parser dependency is not installed."
        presentation = PptxPresentation(str(path))
        lines: list[str] = []
        for slide_index, slide in enumerate(presentation.slides, start=1):
            slide_text: list[str] = []
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        slide_text.append(text)
            if slide_text:
                lines.append(f"Slide {slide_index}: " + " | ".join(slide_text))

        return "\n".join(lines).strip() or "No extractable text found in PPTX."

    def _truncate(self, content: str) -> tuple[str, bool]:
        if len(content) <= self._max_text_chars:
            return content, False
        return content[: self._max_text_chars].rstrip() + "\n... [truncated]", True

    @staticmethod
    def _cell_to_string(value: object) -> str:
        if value is None:
            return ""
        return str(value)
