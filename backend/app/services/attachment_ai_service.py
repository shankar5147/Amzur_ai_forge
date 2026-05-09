from __future__ import annotations

import base64
import csv
import io
import json
import textwrap
from pathlib import Path

import anyio
import openpyxl
from pypdf import PdfReader

try:
    import docx as python_docx  # python-docx
    _HAS_DOCX = True
except ImportError:
    _HAS_DOCX = False

try:
    from pptx import Presentation as PptxPresentation  # python-pptx
    _HAS_PPTX = True
except ImportError:
    _HAS_PPTX = False

from app.models.db_models import Attachment
from app.services.chat_service import ChatService
from app.services.file_storage_service import FileStorageService


class AttachmentAIService:
    def __init__(self, storage: FileStorageService | None = None) -> None:
        self._storage = storage or FileStorageService()

    async def build_context_blocks(
        self,
        attachments: list[Attachment],
        chat_service: ChatService,
    ) -> list[str]:
        blocks: list[str] = []

        for attachment in attachments:
            full_path = self._storage.resolve_path(attachment.file_path)
            if not full_path.exists():
                continue

            if attachment.mime_type.startswith("image/"):
                blocks.append(await self._describe_image(attachment, full_path, chat_service))
                continue

            if attachment.mime_type == "application/pdf":
                blocks.append(await self._extract_pdf_block(attachment, full_path))
                continue

            if attachment.mime_type in {
                "text/csv",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "application/vnd.ms-excel",
            }:
                blocks.append(await self._extract_table_block(attachment, full_path))
                continue

            if attachment.mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                blocks.append(await self._extract_docx_block(attachment, full_path))
                continue

            if attachment.mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
                blocks.append(await self._extract_pptx_block(attachment, full_path))
                continue

            if attachment.mime_type.startswith("video/"):
                blocks.append(
                    f"[Video attachment] {attachment.file_name} ({attachment.mime_type}). "
                    "Mention that video understanding is limited to metadata in this build."
                )
                continue

            # All text-based types (plain, markdown, HTML, RTF, JSON, code, XML, etc.)
            blocks.append(await self._extract_text_block(attachment, full_path))

        return blocks

    async def _describe_image(self, attachment: Attachment, path: Path, chat_service: ChatService) -> str:
        image_bytes = await anyio.to_thread.run_sync(path.read_bytes)
        b64 = base64.b64encode(image_bytes).decode("utf-8")
        image_url = f"data:{attachment.mime_type};base64,{b64}"

        description = await chat_service.analyze_image(
            image_url=image_url,
            filename=attachment.file_name,
        )
        return f"[Image attachment: {attachment.file_name}] {description}"

    async def _extract_pdf_block(self, attachment: Attachment, path: Path) -> str:
        content = await anyio.to_thread.run_sync(self._extract_pdf_text, path)
        return self._wrap_block(
            title=f"PDF attachment: {attachment.file_name}",
            body=content,
            max_chars=5000,
        )

    async def _extract_table_block(self, attachment: Attachment, path: Path) -> str:
        if attachment.mime_type == "text/csv":
            body = await anyio.to_thread.run_sync(self._extract_csv_preview, path)
        else:
            body = await anyio.to_thread.run_sync(self._extract_xlsx_preview, path)

        return self._wrap_block(
            title=f"Table attachment: {attachment.file_name}",
            body=body,
            max_chars=3500,
        )

    async def _extract_docx_block(self, attachment: Attachment, path: Path) -> str:
        body = await anyio.to_thread.run_sync(self._extract_docx_text, path)
        return self._wrap_block(
            title=f"Word document: {attachment.file_name}",
            body=body,
            max_chars=5000,
        )

    async def _extract_pptx_block(self, attachment: Attachment, path: Path) -> str:
        body = await anyio.to_thread.run_sync(self._extract_pptx_text, path)
        return self._wrap_block(
            title=f"PowerPoint presentation: {attachment.file_name}",
            body=body,
            max_chars=5000,
        )

    async def _extract_text_block(self, attachment: Attachment, path: Path) -> str:
        text = await anyio.to_thread.run_sync(path.read_text, "utf-8", "ignore")

        if Path(attachment.file_name).suffix.lower() == ".json":
            try:
                parsed = json.loads(text)
                text = json.dumps(parsed, indent=2)
            except json.JSONDecodeError:
                pass

        return self._wrap_block(
            title=f"Text attachment: {attachment.file_name}",
            body=text,
            max_chars=4000,
        )

    @staticmethod
    def _extract_docx_text(path: Path) -> str:
        if not _HAS_DOCX:
            return "Word document extraction unavailable (python-docx not installed)."
        doc = python_docx.Document(str(path))
        paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
        return "\n".join(paragraphs).strip() or "No extractable text found in document."

    @staticmethod
    def _extract_pptx_text(path: Path) -> str:
        if not _HAS_PPTX:
            return "PowerPoint extraction unavailable (python-pptx not installed)."
        prs = PptxPresentation(str(path))
        slides_text: list[str] = []
        for i, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in para.runs).strip()
                        if line:
                            texts.append(line)
            if texts:
                slides_text.append(f"Slide {i}: " + " | ".join(texts))
        return "\n".join(slides_text).strip() or "No extractable text found in presentation."

    @staticmethod
    def _extract_pdf_text(path: Path) -> str:
        reader = PdfReader(str(path))
        chunks: list[str] = []
        for page in reader.pages:
            chunks.append(page.extract_text() or "")
        return "\n".join(chunks).strip() or "No extractable PDF text found."

    @staticmethod
    def _extract_csv_preview(path: Path) -> str:
        with path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
            rows = list(csv.reader(handle))

        if not rows:
            return "CSV is empty."

        header = rows[0]
        sample_rows = rows[1:11]
        preview = [f"Columns: {', '.join(header)}", f"Rows: {max(len(rows) - 1, 0)}"]
        for idx, row in enumerate(sample_rows, start=1):
            preview.append(f"Row {idx}: {', '.join(row)}")
        return "\n".join(preview)

    @staticmethod
    def _extract_xlsx_preview(path: Path) -> str:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        worksheet = workbook.active

        rows = worksheet.iter_rows(values_only=True)
        first = next(rows, None)
        if first is None:
            workbook.close()
            return "Sheet is empty."

        header = [str(cell) if cell is not None else "" for cell in first]
        preview = [f"Columns: {', '.join(header)}"]

        for idx, row in enumerate(rows, start=1):
            if idx > 10:
                break
            values = [str(cell) if cell is not None else "" for cell in row]
            preview.append(f"Row {idx}: {', '.join(values)}")

        workbook.close()
        return "\n".join(preview)

    @staticmethod
    def _wrap_block(title: str, body: str, max_chars: int) -> str:
        normalized = textwrap.dedent(body).strip() if body else ""
        trimmed = normalized[:max_chars]
        return f"[{title}]\n{trimmed}"
