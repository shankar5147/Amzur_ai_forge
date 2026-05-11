from __future__ import annotations

import csv
import io
import re
from pathlib import Path

import openpyxl
from pypdf import PdfReader

try:
    import docx as python_docx
    _HAS_DOCX = True
except ImportError:  # pragma: no cover
    _HAS_DOCX = False

try:
    from pptx import Presentation as PptxPresentation
    _HAS_PPTX = True
except ImportError:  # pragma: no cover
    _HAS_PPTX = False


class DocumentProcessorService:
    def extract_text(self, file_path: Path, mime_type: str, file_name: str) -> str:
        if mime_type == "application/pdf":
            return self._extract_pdf(file_path)

        if mime_type == "text/csv":
            return self._extract_csv(file_path)

        if mime_type in {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel",
        }:
            return self._extract_xlsx(file_path)

        if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return self._extract_docx(file_path)

        if mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
            return self._extract_pptx(file_path)

        return file_path.read_text("utf-8", errors="ignore")

    def chunk_text(self, text: str, chunk_size: int, overlap: int) -> list[str]:
        normalized = re.sub(r"\r\n?", "\n", text).strip()
        if not normalized:
            return []

        if len(normalized) <= chunk_size:
            return [normalized]

        chunks: list[str] = []
        start = 0
        safe_overlap = max(0, min(overlap, max(0, chunk_size - 1)))

        while start < len(normalized):
            end = min(len(normalized), start + chunk_size)
            segment = normalized[start:end]

            # Prefer splitting at logical boundaries to keep chunk coherence.
            split_at = max(segment.rfind("\n\n"), segment.rfind("\n"), segment.rfind(". "))
            if 0 < split_at < len(segment) - 80:
                segment = segment[: split_at + 1]
                end = start + len(segment)

            cleaned = segment.strip()
            if cleaned:
                chunks.append(cleaned)

            if end >= len(normalized):
                break

            start = max(start + 1, end - safe_overlap)

        return chunks

    @staticmethod
    def _extract_pdf(path: Path) -> str:
        reader = PdfReader(str(path))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        return "\n\n".join(value for value in pages if value)

    @staticmethod
    def _extract_csv(path: Path) -> str:
        with path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
            rows = list(csv.reader(handle))

        rendered: list[str] = []
        for row in rows:
            rendered.append(" | ".join(cell.strip() for cell in row))
        return "\n".join(rendered)

    @staticmethod
    def _extract_xlsx(path: Path) -> str:
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        ws = wb.active
        lines: list[str] = []
        for row in ws.iter_rows(values_only=True):
            values = ["" if value is None else str(value).strip() for value in row]
            lines.append(" | ".join(values))
        wb.close()
        return "\n".join(lines)

    @staticmethod
    def _extract_docx(path: Path) -> str:
        if not _HAS_DOCX:
            return ""
        doc = python_docx.Document(str(path))
        paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text and paragraph.text.strip()]
        return "\n".join(paragraphs)

    @staticmethod
    def _extract_pptx(path: Path) -> str:
        if not _HAS_PPTX:
            return ""
        presentation = PptxPresentation(str(path))
        output = io.StringIO()
        for idx, slide in enumerate(presentation.slides, start=1):
            fragments: list[str] = []
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        fragments.append(text)
            if fragments:
                output.write(f"Slide {idx}: {' | '.join(fragments)}\n")
        return output.getvalue().strip()
